"""Read-only PPTX, notes and PDF check; creates only independent QA evidence/PNGs."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import re
import zipfile
import xml.etree.ElementTree as ET
import pymupdf as fitz
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[2]
meta=json.loads((ROOT/'source/figma/L05_content.json').read_text(encoding='utf-8'))
prs=Presentation(HERE/'editable.pptx');pdf=fitz.open(HERE/'editable.pdf')
figma=fitz.open(HERE/'presentation.pdf')
ns='{http://www.w3.org/2000/svg}'
result={'sha256':{},'slides':len(prs.slides),'size_emu':[prs.slide_width,prs.slide_height],'slide_results':[],'fonts':{},'pdf_fonts':[]}
all_types=Counter();all_fonts=Counter()
for i,slide in enumerate(prs.slides,1):
    expected=meta[i-1]['section']+'\n'+'\n'.join('- '+n for n in meta[i-1]['notes'])+'\nФрагмент лекции: '+meta[i-1]['anchor']
    notes=slide.notes_slide.notes_text_frame.text
    svg_path=ROOT/f'assets/figma/L05/slide-{i:02}.svg'
    root=ET.parse(svg_path).getroot()
    fragments=[''.join(n.itertext()).rstrip('\n') for n in root.iter(ns+'tspan') if ''.join(n.itertext()).rstrip('\n')]
    texts=[s.text for s in slide.shapes if s.has_text_frame and s.text]
    types=Counter(str(s.shape_type) for s in slide.shapes);all_types.update(types)
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if r.text:all_fonts[(r.font.name,bool(r.font.bold))]+=1
    outside=[s.name for s in slide.shapes if min(s.left,s.top)<-1000 or s.left+s.width>prs.slide_width+1000 or s.top+s.height>prs.slide_height+1000]
    page=pdf[i-1]
    norm=lambda s:re.sub(r'\s+','',s)
    pdf_text=norm(page.get_text())
    missing_pdf=[s for s in fragments if norm(s) not in pdf_text]
    spans=[s for b in page.get_text('dict')['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']]
    outside_pdf=[s['text'] for s in spans if s['bbox'][0]<-1 or s['bbox'][1]<-1 or s['bbox'][2]>page.rect.width+1 or s['bbox'][3]>page.rect.height+1]
    row={'slide':i,'shape_types':dict(types),'text_boxes':len(texts),'notes_match':notes==expected,'notes':notes,'missing_svg_fragments':list((Counter(fragments)-Counter(texts)).elements()),'extra_pptx_fragments':list((Counter(texts)-Counter(fragments)).elements()),'pdf_missing_fragments':missing_pdf,'outside_shapes':outside,'outside_pdf':outside_pdf}
    result['slide_results'].append(row)
    for p in [svg_path,HERE/f'slide-{i:02}.pdf',HERE/f'slide-{i:02}.png',HERE/f'editable_slides/slide-{i:02}.png']:
        result['sha256'][str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
out=HERE/'independent_comparison';out.mkdir(exist_ok=True)
for i in range(1,len(prs.slides)+1):
    a=figma[i-1].get_pixmap(matrix=fitz.Matrix(.75,.75),alpha=False)
    b=pdf[i-1].get_pixmap(matrix=fitz.Matrix(1440/pdf[i-1].rect.width,810/pdf[i-1].rect.height),alpha=False)
    im1=Image.frombytes('RGB',[a.width,a.height],a.samples)
    im2=Image.frombytes('RGB',[b.width,b.height],b.samples)
    sheet=Image.new('RGB',(2880,840),'#eeeeee');draw=ImageDraw.Draw(sheet)
    draw.text((16,5),f'{i:02} FIGMA PDF',fill='black',font=ImageFont.load_default(size=20));draw.text((1456,5),f'{i:02} EDITABLE PPTX PDF',fill='black',font=ImageFont.load_default(size=20))
    sheet.paste(im1,(0,30));sheet.paste(im2,(1440,30));sheet.save(out/f'slide-{i:02}.png')
with zipfile.ZipFile(HERE/'editable.pptx') as z:
    result['media_parts']=[n for n in z.namelist() if n.startswith('ppt/media/')]
    result['negative_extents']=[]
    for n in z.namelist():
        if re.fullmatch(r'ppt/slides/slide\d+\.xml',n):
            root=ET.fromstring(z.read(n))
            for el in root.iter():
                if el.tag.endswith('}ext') and any(float(el.attrib.get(k,0))<0 for k in ['cx','cy']):result['negative_extents'].append(n)
result['shape_types']=dict(all_types);result['fonts']=[{'font':k[0],'bold':k[1],'runs':v} for k,v in all_fonts.items()]
result['pdf_fonts']=sorted(set(font[3] for p in pdf for font in p.get_fonts(full=True)))
for p in [HERE/'editable.pptx',HERE/'editable.pdf',HERE/'presentation.pdf',ROOT/'dist/teacher/L05/lecture.pdf',ROOT/'source/lectures/L05.md',ROOT/'source/figma/L05_content.json',ROOT/'source/figma/L05_visuals.json']:
    result['sha256'][str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
(HERE/'powerpoint_review_evidence.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'slides':result['slides'],'shape_types':result['shape_types'],'fonts':result['fonts'],'pdf_fonts':result['pdf_fonts'],'media_parts':result['media_parts'],'notes_mismatch':[r['slide'] for r in result['slide_results'] if not r['notes_match']],'text_mismatch':[r for r in result['slide_results'] if r['missing_svg_fragments'] or r['extra_pptx_fragments'] or r['pdf_missing_fragments']],'outside':[r['slide'] for r in result['slide_results'] if r['outside_shapes'] or r['outside_pdf']],'negative_extents':result['negative_extents']},ensure_ascii=False,indent=2))
