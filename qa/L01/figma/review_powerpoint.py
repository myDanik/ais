from pathlib import Path
from collections import Counter
import json
import hashlib
import zipfile
import unicodedata
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).parent
ROOT = HERE.parents[2]
OUT = HERE / 'powerpoint_review_png'
OUT.mkdir(exist_ok=True)
meta = json.loads((ROOT / 'source/figma/L01_content.json').read_text(encoding='utf-8'))
prs = Presentation(HERE / 'editable.pptx')
slide_results = []
all_types = Counter()
fonts = Counter()
source_ns = '{http://www.w3.org/2000/svg}'
for i, slide in enumerate(prs.slides, 1):
    m = meta[i-1]
    expected_notes = m['section'] + '\n' + '\n'.join('- ' + note for note in m['notes']) + '\nФрагмент лекции: ' + m['anchor']
    actual_notes = slide.notes_slide.notes_text_frame.text
    types = Counter(str(shape.shape_type) for shape in slide.shapes)
    all_types.update(types)
    texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
    bold_texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.text:
                        fonts[(run.font.name, bool(run.font.bold))] += 1
                        if run.font.bold:
                            bold_texts.append(run.text)
    root = ET.parse(HERE / f'slide-{i:02}.svg').getroot()
    svg_fragments = [''.join(t.itertext()).rstrip('\n') for t in root.iter(source_ns+'tspan') if ''.join(t.itertext()).rstrip('\n')]
    # Each exported SVG tspan is a native text shape; compare multiplicities exactly.
    missing = list((Counter(svg_fragments) - Counter(texts)).elements())
    extra = list((Counter(texts) - Counter(svg_fragments)).elements())
    outside = [shape.name for shape in slide.shapes if shape.left < -1000 or shape.top < -1000 or shape.left + shape.width > prs.slide_width + 1000 or shape.top + shape.height > prs.slide_height + 1000]
    slide_results.append({'slide': i, 'shape_types': dict(types), 'text_boxes': len(texts), 'notes_match': actual_notes == expected_notes, 'notes_chars': len(actual_notes), 'missing_svg_fragments': missing, 'extra_pptx_fragments': extra, 'outside': outside, 'bold_texts': bold_texts})
with zipfile.ZipFile(HERE / 'editable.pptx') as package:
    media = [name for name in package.namelist() if name.startswith('ppt/media/')]
    embedded_fonts = [name for name in package.namelist() if name.startswith('ppt/fonts/')]

for start in range(1, len(prs.slides) + 1, 2):
    sheet = Image.new('RGB', (1920, 1160), '#e6e8e7')
    draw = ImageDraw.Draw(sheet)
    for row, number in enumerate(range(start, min(start + 2, len(prs.slides) + 1))):
        y = row * 580
        draw.text((20, y + 8), f'{number:02} FIGMA', fill='black', font=ImageFont.load_default(size=20))
        draw.text((980, y + 8), f'{number:02} POWERPOINT', fill='black', font=ImageFont.load_default(size=20))
        figma = Image.open(HERE / 'review_png' / f'slide-{number:02}.png').convert('RGB')
        pptx = Image.open(HERE / 'editable_slides' / f'slide-{number:02}.png').convert('RGB')
        sheet.paste(figma.resize((960, 540), Image.Resampling.LANCZOS), (0, y + 40))
        sheet.paste(pptx.resize((960, 540), Image.Resampling.LANCZOS), (960, y + 40))
    sheet.save(OUT / f'pair-{(start+1)//2:02}.png')

# Exercise editability on a separate review copy only; the submitted PPTX is unchanged.
probe = Presentation(HERE / 'editable.pptx')
nfc = next(shape for shape in probe.slides[15].shapes if shape.has_text_frame and shape.text == 'NFC')
nfc.text_frame.paragraphs[0].runs[0].text = 'NFD'
circle = next(shape for shape in probe.slides[3].shapes if shape._element.xpath('.//a:prstGeom[@prst="ellipse"]'))
circle_original_left = circle.left
circle.left += 9144
probe.save(HERE / 'review_editability_probe.pptx')
reopened = Presentation(HERE / 'review_editability_probe.pptx')
edit_probe_ok = any(shape.has_text_frame and shape.text == 'NFD' for shape in reopened.slides[15].shapes)
reopened_circle = next(shape for shape in reopened.slides[3].shapes if shape._element.xpath('.//a:prstGeom[@prst="ellipse"]'))
geometry_probe_ok = reopened_circle.left == circle_original_left + 9144
result = {'sha256': hashlib.sha256((HERE / 'editable.pptx').read_bytes()).hexdigest(), 'slides': len(prs.slides), 'size_emu': [prs.slide_width, prs.slide_height], 'shape_types': dict(all_types), 'fonts': [{'name': name, 'bold': bold, 'runs': count} for (name, bold), count in fonts.items()], 'media_parts': media, 'embedded_font_parts': embedded_fonts, 'edit_probe_ok': edit_probe_ok, 'geometry_probe_ok': geometry_probe_ok, 'slide_results': slide_results}
(HERE / 'powerpoint_review_evidence.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({key: result[key] for key in ['slides', 'shape_types', 'media_parts', 'embedded_font_parts', 'edit_probe_ok']}, ensure_ascii=True))
print(json.dumps({'notes_mismatch': [r['slide'] for r in slide_results if not r['notes_match']], 'text_mismatch': [r['slide'] for r in slide_results if r['missing_svg_fragments'] or r['extra_pptx_fragments']], 'outside': [r['slide'] for r in slide_results if r['outside']]}, ensure_ascii=True))
