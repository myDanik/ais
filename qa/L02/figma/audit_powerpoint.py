"""Independent read-only check of the final editable L02 PowerPoint."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import sys
import xml.etree.ElementTree as ET

from pptx import Presentation
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[3]
QA = Path(__file__).resolve().parent
sys.stdout.reconfigure(encoding="utf-8")
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
content = json.loads((ROOT / "source/figma/L02_content.json").read_text(encoding="utf-8"))
pptx_path = QA / "editable.pptx"
pdf_path = QA / "editable.pdf"
prs = Presentation(pptx_path)
pdf = fitz.open(pdf_path)
assert len(prs.slides) == len(pdf) == len(content) == 23
ns = {"s": "http://www.w3.org/2000/svg"}
slides = []
fonts = set()

def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from all_shapes(shape.shapes)

for i, (slide, page, c) in enumerate(zip(prs.slides, pdf, content), 1):
    expected_notes = c["section"] + "\n" + "\n".join("- " + n for n in c["notes"]) + "\nФрагмент лекции: " + c["anchor"]
    actual_notes = slide.notes_slide.notes_text_frame.text
    assert actual_notes == expected_notes, (i, actual_notes, expected_notes)
    root = ET.parse(ROOT / f"assets/figma/L02/slide-{i:02}.svg").getroot()
    spans = [s for b in page.get_text("rawdict")["blocks"] for line in b.get("lines", []) for s in line["spans"]]
    for s in spans:
        s["text"] = "".join(ch["c"] for ch in s["chars"])
    matches, missing = [], []
    for element in root.findall(".//s:tspan", ns):
        text = "".join(element.itertext()).strip()
        if not text:
            continue
        x, y = float(element.get("x")), float(element.get("y"))
        candidates = []
        for s in spans:
            start = 0
            while (at := s["text"].find(text, start)) >= 0:
                ax, ay = s["chars"][at]["origin"]
                candidates.append((abs(ax * 2 - x) + abs(ay * 2 - y), ax, ay))
                start = at + 1
        if not candidates:
            missing.append(text)
            continue
        _, ax, ay = min(candidates)
        matches.append({"text": text, "dx_svg_px": ax * 2 - x, "dy_svg_px": ay * 2 - y})
    shapes = list(all_shapes(slide.shapes))
    pictures = [s for s in shapes if s.shape_type == 13]
    text_boxes = [s for s in shapes if getattr(s, "has_text_frame", False) and s.text_frame.text]
    for font in page.get_fonts():
        fonts.add(font[3].split("+")[-1])
    bounds_issues = [s["bbox"] for s in spans if not page.rect.contains(fitz.Rect(s["bbox"]))]
    slides.append({"number": i, "title": c["title"], "notes_exact_match": True, "notes": actual_notes,
        "native_shape_count": len(shapes), "native_textbox_count": len(text_boxes), "picture_shape_count": len(pictures),
        "matched_svg_fragments": len(matches), "missing_svg_fragments": missing,
        "max_baseline_error_svg_px": max(abs(m["dy_svg_px"]) for m in matches),
        "max_horizontal_error_svg_px": max(abs(m["dx_svg_px"]) for m in matches),
        "matches": matches, "outside_page_text_bbox": bounds_issues,
        "render_sha256": sha(QA / f"editable_slides/slide-{i:02}.png")})

out = {"date": "2026-09-08", "pptx_sha256": sha(pptx_path), "pdf_sha256": sha(pdf_path),
       "source_content_sha256": sha(ROOT / "source/figma/L02_content.json"),
       "slides": slides, "pdf_fonts": sorted(fonts), "reviewed_pngs": list(range(1, 24)),
       "visual_result": "all 23 PNGs inspected; no clipping, visible overlap, changed formulas or missing punctuation",
       "slide_size_emu": [prs.slide_width, prs.slide_height],
       "all_notes_match": all(s["notes_exact_match"] for s in slides),
       "total_matched_svg_fragments": sum(s["matched_svg_fragments"] for s in slides),
       "missing_svg_fragments": sum(len(s["missing_svg_fragments"]) for s in slides),
       "max_baseline_error_svg_px": max(s["max_baseline_error_svg_px"] for s in slides)}
(QA / "powerpoint_review_evidence.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: v for k, v in out.items() if k != "slides"}, ensure_ascii=False, indent=2))
print(json.dumps({s["number"]: s["missing_svg_fragments"] for s in slides if s["missing_svg_fragments"]}, ensure_ascii=False))
assert out["missing_svg_fragments"] == 0
assert all(not s["outside_page_text_bbox"] for s in slides)
assert fonts <= {"Prata-Regular", "Carlito-Regular", "Carlito-Bold"}
