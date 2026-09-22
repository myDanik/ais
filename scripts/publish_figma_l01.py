"""Publish the reviewed L01 files at the existing course paths."""
import hashlib,json,shutil
from pathlib import Path
import pymupdf as fitz
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
Q=ROOT/'qa/L01/figma'
T=ROOT/'dist/teacher/L01'
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def main():
    current=json.loads((ROOT/'source/figma/L01_content.json').read_text(encoding='utf8'))
    candidate=Presentation(Q/'editable.pptx')
    assert len(candidate.slides)==len(current),'PPTX count differs from current lecture'
    for s,m in zip(candidate.slides,current):
        notes=s.notes_slide.notes_text_frame.text
        assert m['anchor'] in notes and '\n'.join('- '+n for n in m['notes']) in notes,'Reviewed snapshot has outdated lecture notes; rebuild and review it first'
    for name in ['presentation_review.md','powerpoint_review.md','language_review.md']:
        text=(Q/name).read_text(encoding='utf8')
        assert 'ОК' in text,name+' has not been accepted'
        if name=='powerpoint_review.md':
            assert digest(Q/'editable.pptx') in text,'Reviewer has not checked this PPTX'
            assert 'Промежуточный итог' not in text[:700],'Final PPTX review is still pending'
    for ext in ['pptx','pdf']:
        shutil.copy2(Q/f'editable.{ext}',T/f'presentation.{ext}')
    preview=(T/'presentation_slides').resolve()
    assert preview.is_relative_to(ROOT) and preview==ROOT/'dist/teacher/L01/presentation_slides'
    preview.mkdir(exist_ok=True)
    # Only replace known preview files; never remove other contents recursively.
    for p in preview.iterdir():
        if p.is_file() and p.suffix.lower() in {'.png','.jpg'}:p.unlink()
    for p in (Q/'editable_slides').iterdir():
        if p.is_file() and p.suffix.lower() in {'.png','.jpg'}:shutil.copy2(p,preview/p.name)
    qa=json.loads((Q/'editable.qa.json').read_text(encoding='utf8'))
    qa['file']=str(T/'presentation.pptx');qa['visual_review_required']=False
    qa['visual_review']='qa/L01/figma/powerpoint_review.md'
    (T/'presentation.qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    for audience in ['student','teacher']:
        folder=ROOT/f'dist/{audience}/L01'
        for ext in ['pptx','pdf']:
            target=folder/f'presentation.{ext}'
            if target!=T/f'presentation.{ext}':shutil.copy2(T/f'presentation.{ext}',target)
        (folder/'Figma.url').write_text('[InternetShortcut]\nURL=https://www.figma.com/slides/OBqHFG8UVRENgZheAnNGMG\n',encoding='ascii')
    reviews=ROOT/'dist/teacher/reviews'
    reviews.mkdir(exist_ok=True)
    for name in ['presentation_review','powerpoint_review','language_review']:
        shutil.copy2(Q/f'{name}.md',reviews/f'L01_figma_{name}.md')
    p=ROOT/'source/figma/L01_file.json';data=json.loads(p.read_text(encoding='utf8'))
    data['status']='approved';data['reviewed']='2026-09-08'
    data['deliverables']={'pptx':'dist/teacher/L01/presentation.pptx','pdf':'dist/teacher/L01/presentation.pdf'}
    data['prepared']='23 native Figma slides; editable PPTX derived from exported SVG geometry; independent visual and language reviews passed'
    data['sha256']={ext:digest(T/f'presentation.{ext}') for ext in ['pptx','pdf']}
    p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps({'published':str(T),'slides':23,'sha256':data['sha256']},ensure_ascii=False))

if __name__=='__main__':main()
