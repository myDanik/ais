"""Render the actual PPTX with LibreOffice and Poppler, then report structural/text defects."""
import argparse,json,os,shutil,subprocess,sys,zipfile
from pathlib import Path
import lxml.etree as ET
import pymupdf as fitz
from PIL import Image,ImageDraw
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from build_materials import to_pdf
def render(path):
    path=path.resolve();to_pdf(path)
    out=path.parent/(path.stem+'_slides');out.mkdir(exist_ok=True)
    exe=next((ROOT/'.tools/poppler').glob('*/Library/bin/pdftoppm.exe'),None)
    with fitz.open(path.with_suffix('.pdf')) as pdf:
        if exe:
            subprocess.run([str(exe),'-png','-r','144',str(path.with_suffix('.pdf')),str(out/'slide')],check=True,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        else:
            for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(out/f'slide-{i+1}.png')
        fonts=sorted({f[3] for page in pdf for f in page.get_fonts()})
        errors=[]
        for i,page in enumerate(pdf):
            if not page.get_text().strip():errors.append(f'Page {i+1}: no extractable text')
            for block in page.get_text('dict')['blocks']:
                for ln in block.get('lines',[]):
                    for span in ln['spans']:
                        x0,y0,x1,y1=span['bbox']
                        if span['text'].strip() and (x0<0 or y0<0 or x1>page.rect.width+.5 or y1>page.rect.height+.5):errors.append(f'Page {i+1}: text outside slide: {span["text"][:60]}')
    prs=Presentation(path)
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
                for e in ET.fromstring(z.read(name)).iter():
                    if e.tag.endswith('}ext') and any(int(e.get(k,'0'))<0 for k in ['cx','cy']):errors.append('Negative extent in '+name)
    for i,slide in enumerate(prs.slides):
        if not slide.has_notes_slide:errors.append(f'Slide {i+1}: missing notes')
        for sh in slide.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+2000 or sh.top+sh.height>prs.slide_height+2000:errors.append(f'Slide {i+1}: object outside slide')
    images=sorted(out.glob('slide-*.png'),key=lambda p:int(p.stem.split('-')[-1]))
    for start in range(0,len(images),6):
        sheet=Image.new('RGB',(1440,1350),'#E5E7E7');draw=ImageDraw.Draw(sheet)
        for j,p in enumerate(images[start:start+6]):
            im=Image.open(p);im.thumbnail((690,389));x=20+j%2*720;y=20+j//2*450;sheet.paste(im,(x,y));draw.text((x,y+400),str(start+j+1),fill='black')
        sheet.save(out/f'contact_{start//6+1:02}.jpg',quality=94)
    report={'file':str(path),'slides':len(prs.slides),'renderer':'LibreOffice + Poppler' if exe else 'LibreOffice + PyMuPDF','fonts':fonts,'errors':errors,'visual_review_required':True}
    (path.parent/(path.stem+'.qa.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))
    if errors:raise RuntimeError('Fix the reported defects before visual review')
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('file',type=Path);args=ap.parse_args();render(args.file)
