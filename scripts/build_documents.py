"""Readable lecture manuscripts and lab handouts with real headings and links."""
import argparse,re,sys,shutil
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from build_materials import to_pdf

def hyperlink(p,label,url):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'6D2E46');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline)
    run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)
def inline(p,text):
    chunks=re.split(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*.+?\*\*)',text)
    for part in chunks:
        match=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if match:hyperlink(p,match[1],match[2]);continue
        code=part.startswith('`') and part.endswith('`')
        bold=part.startswith('**') and part.endswith('**')
        run=p.add_run(part[1:-1] if code else part[2:-2] if bold else part)
        if code:run.font.name='Courier New';run.font.size=Pt(10.5)
        if bold:run.bold=True
def convert(src,dst):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.15);sec.bottom_margin=Cm(2.1)
    sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
    sec.header_distance=Cm(.95);sec.footer_distance=Cm(1.0)
    normal=doc.styles['Normal'];normal.font.name='Cambria';normal.font.size=Pt(12)
    normal.font.color.rgb=RGBColor.from_string('1B1D22')
    normal.paragraph_format.line_spacing=1.17
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size,font in [('Title',30,'Cambria'),('Heading 1',19,'Cambria'),('Heading 2',14,'Arial')]:
        st=doc.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.bold=name!='Title'
        st.font.color.rgb=RGBColor.from_string('171A20')
        st.paragraph_format.space_before=Pt(18 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(9);st.paragraph_format.keep_with_next=True
        for border in st.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');normal.element.get_or_add_rPr().append(lang)
    header=sec.header.paragraphs[0];header.text='Системы искусственного интеллекта · NLP'
    for r in header.runs:r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('717782')
    footer=sec.footer.paragraphs[0];footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    lines=src.read_text(encoding='utf-8').splitlines();i=0;first=True;subtitle=False
    while i<len(lines):
        line=lines[i]
        if not line.strip() or line.startswith('<!--'):i+=1;continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:],'Title');subtitle=src.parent.name=='lectures';i+=1;continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:],'Heading 1');subtitle=False;i+=1;continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:],'Heading 2');i+=1;continue
        if line.startswith('```'):
            body=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):body.append(lines[i]);i+=1
            p=doc.add_paragraph();p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(12)
            p.paragraph_format.left_indent=Cm(.25)
            run=p.add_run('\n'.join(body));run.font.name='Courier New';run.font.size=Pt(10.5)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F3F4F6');p._p.get_or_add_pPr().append(shade)
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                parts=[x.strip().replace(r'\|','|') for x in re.split(r'(?<!\\)\|',lines[i].strip().strip('|'))]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in parts):rows.append(parts)
                i+=1
            table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
            weights={2:[4,12.4],3:[3.0,5.7,7.7]}.get(len(rows[0]),[16.4/len(rows[0])]*len(rows[0]))
            if src.parent.name=='lectures' and src.stem=='L02' and len(rows[0])==3:weights=[8.4,4.0,4.0]
            if src.parent.name=='plans' and src.stem=='L07' and len(rows[0])==3:weights=[4.0,5.2,7.2]
            for col,width in zip(table.columns,weights):col.width=Cm(width)
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                for cell,width,text in zip(cells,weights,row):
                    cell.width=Cm(width);p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
                    if src.parent.name=='lectures' and src.stem in {'L01','L02'} and len(rows)<=6:
                        p.paragraph_format.keep_with_next=ri<len(rows)-1
                    inline(p,text)
                    for run in p.runs:run.font.name='Arial';run.font.size=Pt(10.5);run.bold=ri==0
                    tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom')
                    b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6' if ri==0 else '3');b.set(qn('w:color'),'A9AEB8' if ri==0 else 'E0E3E8');borders.append(b);tcpr.append(borders)
                table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if ri==0:table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        p=doc.add_paragraph()
        if src.stem=='slide_map' and line.startswith('Раздел:'):
            p.paragraph_format.keep_with_next=True
        if line.startswith('- '):
            if src.stem=='speaker_notes':
                following=next((x for x in lines[i+1:] if x.strip()),'')
                p.paragraph_format.keep_with_next=following.startswith('- ')
            p.paragraph_format.left_indent=Cm(.35);p.paragraph_format.first_line_indent=Cm(-.35);line='– '+line[2:]
        inline(p,line)
        if subtitle:
            for r in p.runs:r.font.size=Pt(15);r.italic=True;r.font.color.rgb=RGBColor.from_string('626773')
            p.paragraph_format.space_after=Pt(17);subtitle=False
        i+=1
    dst.parent.mkdir(parents=True,exist_ok=True);doc.save(dst)
def build(n):
    lid=f'L{n:02}'
    sources={'lecture':ROOT/'source/lectures'/f'{lid}.md','plan':ROOT/'source/plans'/f'{lid}.md','lab_guide':ROOT/'source/labs'/f'{lid}_guide.md'}
    for stem,src in sources.items():
        if not src.exists():continue
        out=ROOT/'dist/teacher'/lid
        convert(src,out/f'{stem}.docx');shutil.copy2(src,out/f'{stem}.md');to_pdf(out/f'{stem}.docx')
        if stem!='plan':
            dst=ROOT/'dist/student'/lid;dst.mkdir(parents=True,exist_ok=True)
            for ext in ['md','docx','pdf']:shutil.copy2(out/f'{stem}.{ext}',dst/f'{stem}.{ext}')
    print(f'Built revised documents {lid}',flush=True)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--lesson',type=int,required=True);args=ap.parse_args();build(args.lesson)
