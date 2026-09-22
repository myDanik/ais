"""Check portable contents, actual output provenance, slide links and source sync."""
import hashlib,json,os,re,sys,zipfile
from pathlib import Path
import nbformat,pymupdf
from pptx import Presentation
R=Path(__file__).resolve().parents[1];D=R/'dist';Q=R/'qa'
sys.path.insert(0,str(R/'scripts'))
from provenance import fingerprint
os.environ['NLP_MODELS_DIR']=str(R/'.cache/models')
errors=[];lessons=[];pdf_count=0
def check(ok,message):
 if not ok:errors.append(message)
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
for n in range(1,9):
 lid=f'L{n:02}';teacher=D/'teacher'/lid;student=D/'student'/lid
 spec=json.loads((R/'source/slides'/f'{lid}.json').read_text(encoding='utf-8'))
 lecture=(R/spec['lecture']).read_text(encoding='utf-8');prs=Presentation(teacher/'presentation.pptx')
 if (R/'source/figma'/f'{lid}_file.json').exists():
  figma_meta=json.loads((R/'source/figma'/f'{lid}_file.json').read_text(encoding='utf-8'))
  check(figma_meta.get('status')=='approved',lid+' Figma not approved')
  for ext,sha in figma_meta.get('sha256',{}).items():
   check(digest(teacher/f'presentation.{ext}')==sha,lid+' published '+ext+' differs from accepted file')
  cues=re.findall(r'^### Слайд (\d+)\. (.+)$',lecture,re.M)
  check([int(i) for i,_ in cues]==list(range(1,len(spec['slides'])+1)),lid+' visible lecture slide cues')
  with pymupdf.open(teacher/'lecture.pdf') as manuscript_pdf:
   printed=re.sub(r'\s+','', ''.join(page.get_text(sort=True) for page in manuscript_pdf))
 check(len(prs.slides)==len(spec['slides']),lid+' slide count')
 check(digest(teacher/'presentation.pptx')==digest(student/'presentation.pptx'),lid+' student PPTX differs')
 check(digest(teacher/'presentation.pdf')==digest(student/'presentation.pdf'),lid+' student PDF differs')
 check(lecture==(teacher/'lecture.md').read_text(encoding='utf-8'),lid+' manuscript out of date')
 for i,(slide,s) in enumerate(zip(prs.slides,spec['slides']),1):
  marker='<!-- '+s['section']+' -->'
  check(marker in lecture and s['anchor'] in lecture.split(marker,1)[-1].split('\n## ',1)[0],f'{lid}/{i} anchor')
  check(slide.has_notes_slide and s['notes'] in slide.notes_slide.notes_text_frame.text,f'{lid}/{i} notes')
  if (R/'source/figma'/f'{lid}_file.json').exists():
   cue=f'### Слайд {i}. {s["title"]}'
   check(cue in lecture,f'{lid}/{i} manuscript heading differs from slide')
   block=lecture.split(cue,1)[-1].split('\n### Слайд ',1)[0].strip()
   check(block.startswith(s['anchor']),f'{lid}/{i} slide switch does not match passage start')
   check(re.sub(r'\s+','',cue[4:]) in printed,f'{lid}/{i} slide cue absent from lecture PDF')
 teacher_nb=nbformat.read(teacher/'solution.ipynb',as_version=4);student_nb=nbformat.read(student/'lab.ipynb',as_version=4)
 nbformat.validate(teacher_nb);nbformat.validate(student_nb)
 code=[c for c in teacher_nb.cells if c.cell_type=='code']
 check(all(c.execution_count for c in code),lid+' missing execution counts')
 check(not any(o.output_type=='error' for c in code for o in c.outputs),lid+' error output')
 check(teacher_nb.metadata.get('execution_provenance')==fingerprint(D/'teacher',teacher_nb),lid+' provenance mismatch after staging')
 check(not any('solution' in c.metadata.get('tags',[]) for c in student_nb.cells),lid+' solution tag in student notebook')
 check(all(not c.outputs and c.execution_count is None for c in student_nb.cells if c.cell_type=='code'),lid+' student stored outputs')
 for stem in ['lecture','plan','lab_guide','speaker_notes','slide_map']:
  for ext in ['md','docx','pdf']:check((teacher/f'{stem}.{ext}').exists(),lid+' missing '+stem+'.'+ext)
 log=json.loads((Q/lid/'mcp_calls.json').read_text(encoding='utf-8'))
 calls=log if isinstance(log,list) else log.get('operations',[])
 check(bool(calls) and not any(c.get('result',{}).get('isError') for c in calls),lid+' MCP operation log')
 lessons.append({'lesson':lid,'slides':len(prs.slides),'executed_cells':len(code),'pptx_sha256':digest(teacher/'presentation.pptx'),'mcp_calls':len(calls)})
for audience in ['teacher','student']:
 base=D/audience
 for p in base.rglob('*.pdf'):
  with pymupdf.open(p) as pdf:
   check(len(pdf)>0,str(p)+' empty PDF');pdf_count+=1
   for i,page in enumerate(pdf):
    for block in page.get_text('dict')['blocks']:
     for ln in block.get('lines',[]):
      for span in ln['spans']:
       x0,y0,x1,y1=span['bbox']
       check(not span['text'].strip() or (x0>=-.5 and y0>=-.5 and x1<=page.rect.width+.5 and y1<=page.rect.height+.5),f'{p.relative_to(D)} page {i+1}: out-of-page text {span["text"][:40]}')
 for p in base.rglob('*.html'):
  for href in re.findall(r'href="([^"]+)"',p.read_text(encoding='utf-8')):
   if not re.match(r'^[a-z]+:',href):check((p.parent/href).exists(),f'{p}: missing link {href}')
check(not any(p.name.startswith('solution') or 'EXAM_KEYS' in p.name or p.name=='run_records.json' for p in (D/'student').rglob('*')), 'Teacher-only file in student distribution')
check(len(json.loads((R/'course/exam_bank.json').read_text(encoding='utf-8')))==16,'Exam bank count')
report={'status':'passed' if not errors else 'failed','slides':sum(x['slides'] for x in lessons),'executed_cells':sum(x['executed_cells'] for x in lessons),'pdf_files_checked':pdf_count,'lessons':lessons,'errors':errors}
(Q/'release_integrity.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
if errors:raise SystemExit(1)
