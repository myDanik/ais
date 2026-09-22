"""Check the actual lecture PDF, slide cues and unchanged slide artwork."""
import hashlib,json,re,zipfile
from pathlib import Path
import pymupdf as fitz
from PIL import Image,ImageDraw
from pptx import Presentation

R=Path(__file__).resolve().parents[1]
Q=R/'qa/L01/figma';Q.mkdir(parents=True,exist_ok=True)
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def compact(s):return re.sub(r'\s+','',s)
lecture=(R/'source/lectures/L01.md').read_text(encoding='utf8')
meta=json.loads((R/'source/figma/L01_content.json').read_text(encoding='utf8'))
headings=list(re.finditer(r'^### Слайд (\d+)\. (.+)$',lecture,re.M))
assert [int(m[1]) for m in headings]==list(range(1,24))
pdf=fitz.open(R/'dist/teacher/L01/lecture.pdf')
pages=[compact(p.get_text(sort=True)) for p in pdf]
errors=[];locations=[]
for i,(m,s) in enumerate(zip(headings,meta),1):
    assert m[2]==s['title']
    block=lecture[m.end():headings[i].start() if i<len(headings) else len(lecture)]
    assert s['anchor'] in block,(i,s['anchor'])
    heading=f'Слайд {i}. {s["title"]}'
    matches=[j+1 for j,p in enumerate(pages) if compact(heading) in p]
    assert len(matches)==1,(i,matches)
    locations.append({'slide':i,'title':s['title'],'lecture_page':matches[0],'anchor':s['anchor']})
    assert block.strip().startswith(s['anchor']),(i,'anchor is not the start of the spoken passage')
for i,p in enumerate(pdf,1):
    for b in p.get_text('dict')['blocks']:
        for line in b.get('lines',[]):
            for s in line['spans']:
                x0,y0,x1,y1=s['bbox']
                if s['text'].strip() and (x0<0 or y0<0 or x1>p.rect.width+.5 or y1>p.rect.height+.5):errors.append(f'Page {i}: out-of-page text')
    out=Q/f'lecture-page-{i:02}.png'
    p.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False).save(out)
images=sorted(Q.glob('lecture-page-*.png'))
for start in range(0,len(images),4):
    sheet=Image.new('RGB',(1240,1800),'#dddddd');draw=ImageDraw.Draw(sheet)
    for j,p in enumerate(images[start:start+4]):
        im=Image.open(p);im.thumbnail((590,835));x=20+(j%2)*620;y=20+(j//2)*900
        sheet.paste(im,(x,y));draw.text((x,y+843),f'Page {start+j+1}',fill='black')
    sheet.save(Q/f'lecture-contact-{start//4+1}.jpg',quality=94)
old=R/'.cache/L01_before_lecture_sync.pptx';new=R/'dist/teacher/L01/presentation.pptx'
with zipfile.ZipFile(old) as a,zipfile.ZipFile(new) as b:
    visual=[n for n in a.namelist() if n.startswith(('ppt/slides/','ppt/media/','ppt/theme/','ppt/slideMasters/','ppt/slideLayouts/'))]
    changed=[n for n in visual if a.read(n)!=b.read(n)]
    assert not changed,changed
prs=Presentation(new)
for i,(s,m) in enumerate(zip(prs.slides,meta),1):
    notes=s.notes_slide.notes_text_frame.text
    assert m['anchor'] in notes,(i,'PPTX anchor')
    assert '\n'.join('- '+n for n in m['notes']) in notes,(i,'PPTX notes')
report={'status':'passed' if not errors else 'failed','lecture_pages':len(pdf),'slides':len(meta),
        'slide_visual_parts_compared':len(visual),'changed_visual_parts':changed,
        'lecture_pdf_sha256':digest(R/'dist/teacher/L01/lecture.pdf'),'pptx_sha256':digest(new),
        'locations':locations,'errors':errors}
(Q/'lecture_sync_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({k:v for k,v in report.items() if k!='locations'},ensure_ascii=False))
assert not errors
