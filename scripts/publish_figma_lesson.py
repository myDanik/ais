"""Publish one independently accepted Figma lecture at the canonical paths."""
import argparse,hashlib,json,re,shutil,sys
from pathlib import Path
import pymupdf as fitz
from pptx import Presentation
R=Path(__file__).resolve().parents[1]
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main(lid):
 assert re.fullmatch(r'L0[1-8]',lid)
 q=R/f'qa/{lid}/figma';teacher=R/f'dist/teacher/{lid}';student=R/f'dist/student/{lid}'
 for name in ['language_review.md','presentation_review.md','powerpoint_review.md']:
  report=(q/name).read_text(encoding='utf8');assert 'ОК' in report[:900],name
  assert not any(x in report[:600].lower() for x in ['ожидается','предварительный','требуются']),name
  if name=='powerpoint_review.md':assert digest(q/'editable.pptx') in report,name+' hash'
 meta=json.loads((R/f'source/figma/{lid}_content.json').read_text(encoding='utf8'))
 prs=Presentation(q/'editable.pptx');assert len(prs.slides)==len(meta)
 for s,m in zip(prs.slides,meta):
  notes=s.notes_slide.notes_text_frame.text;assert m['anchor'] in notes and '\n'.join('- '+x for x in m['notes']) in notes
 for e in ['pptx','pdf']:shutil.copy2(q/f'editable.{e}',teacher/f'presentation.{e}')
 for e in ['md','docx','pdf']:
  if e=='md':shutil.copy2(R/f'source/lectures/{lid}.md',teacher/'lecture.md')
  shutil.copy2(teacher/f'lecture.{e}',student/f'lecture.{e}')
 preview=(teacher/'presentation_slides').resolve();assert preview.is_relative_to(R) and preview==R/f'dist/teacher/{lid}/presentation_slides'
 preview.mkdir(exist_ok=True)
 for p in preview.iterdir():
  if p.is_file() and p.suffix.lower() in {'.png','.jpg'}:p.unlink()
 for p in (q/'editable_slides').iterdir():
  if p.is_file():shutil.copy2(p,preview/p.name)
 qa=json.loads((q/'editable.qa.json').read_text(encoding='utf8'));qa.update(file=str(teacher/'presentation.pptx'),visual_review_required=False,visual_review=f'qa/{lid}/figma/powerpoint_review.md')
 (teacher/'presentation.qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf8')
 file_path=R/f'source/figma/{lid}_file.json';file=json.loads(file_path.read_text(encoding='utf8'))
 for i,s in enumerate(file['slides']):s['title']=meta[i]['title']
 for audience in [teacher,student]:(audience/'Figma.url').write_text('[InternetShortcut]\nURL='+file['url']+'\n',encoding='ascii')
 for name in ['language_review','presentation_review','powerpoint_review','facts_review']:
  if (q/f'{name}.md').exists():shutil.copy2(q/f'{name}.md',R/f'dist/teacher/reviews/{lid}_figma_{name}.md')
 with fitz.open(teacher/'lecture.pdf') as pdf:pages=len(pdf)
 file.update(status='approved',reviewed='2026-09-21',sha256={e:digest(teacher/f'presentation.{e}') for e in ['pptx','pdf']},lecture_sync={'status':'approved','pages':pages,'visible_slide_cues':len(meta),'pdf_sha256':digest(teacher/'lecture.pdf')})
 file_path.write_text(json.dumps(file,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 spec_path=R/f'source/slides/{lid}.json';spec=json.loads(spec_path.read_text(encoding='utf8'));spec['output']=f'dist/teacher/{lid}/presentation.pptx';spec_path.write_text(json.dumps(spec,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 n=int(lid[1:]);code="import fs from 'node:fs/promises';\nimport path from 'node:path';\nimport {spawnSync} from 'node:child_process';\nimport type {Deck} from './theme.js';\n"
 code+=f"export async function buildLecture{n}():Promise<Deck>{{\n const r=spawnSync(path.join(process.cwd(),'.venv/Scripts/python.exe'),['-X','utf8','scripts/figma_to_pptx_spec.py','--lesson','{lid}'],{{cwd:process.cwd(),stdio:'inherit',windowsHide:true}});\n if(r.status!==0)throw new Error('Figma snapshot conversion failed');\n return JSON.parse(await fs.readFile('source/slides/{lid}.json','utf8')) as Deck;\n}}\n"
 (R/f'src/lecture{n}.ts').write_text(code,encoding='utf8')
 legacy=R/f'source/slides/{lid}_notes.json'
 if legacy.exists():legacy.unlink()
 sys.path.insert(0,str(R/'scripts'));from package_course import sidecars;sidecars([n])
 for e in ['pptx','pdf']:shutil.copy2(teacher/f'presentation.{e}',student/f'presentation.{e}')
 print(json.dumps({'published':lid,'slides':len(meta),'lecture_pages':pages,'pptx_sha256':digest(teacher/'presentation.pptx')},ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('lesson');a=ap.parse_args();main(a.lesson)
