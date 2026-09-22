"""Render the current course documents."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Cm
R=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(R/'scripts'))
from build_documents import convert,to_pdf
C=R/'course'
assert (C/'RPD.md').exists(), 'The current manuscript is missing'
for path in sorted(C.glob('*.md')):
 if path.stem=='REBUILD':continue
 dest=path.with_suffix('.docx');convert(path,dest);doc=Document(dest)
 for p in doc.paragraphs:
  if path.stem in {'EXAM_TICKETS','EXAM_KEYS'} and p.style.name=='Heading 1' and p.text.startswith('Билет '):p.paragraph_format.page_break_before=True
 for table in doc.tables:
  headers=[c.text for c in table.rows[0].cells];n=len(headers);weights=None
  if path.stem=='ASSESSMENT' and n==5:weights=[2.4,.9,4.25,4.6,4.25]
  elif headers[0]=='Код':weights=[1.0,8.0,7.4]
  elif headers[0]=='Тема':weights=[.8,.9,14.7]
  elif n==6:weights=[.6,4.9,.9,.9,1.2,7.9]
  elif path.stem=='RPD' and n==2:weights=[13.9,2.5]
  if weights:
   for col,w in zip(table.columns,weights):col.width=Cm(w)
   for row in table.rows:
    for cell,w in zip(row.cells,weights):cell.width=Cm(w)
 doc.save(dest);to_pdf(dest);print('Built common document',path.stem,flush=True)
